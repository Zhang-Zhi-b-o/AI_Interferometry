"""三个性能指标的中文呈现及 DOCX/JSON 导出。"""
from __future__ import annotations

from datetime import datetime
import json
from pathlib import Path
from uuid import uuid4

from app.services.performance import MODES, summarize_turns


def rate_text(value):
    return "未评定" if value is None else f"{value:.1%}"


def seconds_text(value):
    return "未记录" if value is None else f"{value:.3f} s"


def summary_rows(snapshot: dict) -> list[list[str]]:
    retrieval = snapshot.get("retrieval")
    rows = [["指标 / 范围", "结果", "样本说明"]]
    if retrieval:
        rows.append(["检索命中率 Hit@3", rate_text(retrieval["hit_at_3"]),
                     f"{retrieval['hits']}/{retrieval['count']} 题；无效标注 {retrieval['invalid_labels']}；查询错误 {retrieval['errors']}"])
        rows.append(["检索响应耗时（中位数 / P95）",
                     seconds_text(retrieval["latency"]["median_s"]) + " / " + seconds_text(retrieval["latency"]["p95_s"]),
                     f"本地固定题集，{retrieval['latency']['count']} 次查询"])
    else:
        rows.append(["检索命中率 Hit@3", "未评定", "请先运行检索评测"])
    groups = summarize_turns(snapshot.get("turns", []))
    if not groups:
        rows.append(["任务完成率 / 助手响应耗时", "未评定 / 未记录", "尚无助手轮次，不以软件测试通过率替代"])
    for group in groups:
        scope = MODES[group["mode"]] + " · " + ("模拟数据" if group["data_source"] == "demo" else "现场数据")
        if group["model"]:
            scope += " · " + group["model"]
        rows.append(["任务完成率：" + scope, rate_text(group["task_rate"]),
                     f"完成 {group['successes']}/已验收 {group['judged']}；待验收 {group['pending']}；取消 {group['cancelled']}"])
        rows.append(["助手响应耗时：" + scope,
                     seconds_text(group["latency"]["median_s"]) + " / " + seconds_text(group["latency"]["p95_s"]),
                     f"中位数 / P95；{group['latency']['count']} 轮，含异常，排除用户取消"])
    return rows


def export_performance(snapshot: dict, directory: str | Path) -> dict:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt

    from app.services.report import create_document

    data = json.loads(json.dumps(snapshot, ensure_ascii=False, allow_nan=False))
    data["generated_at"] = datetime.now().astimezone().isoformat(timespec="seconds")
    data["groups"] = summarize_turns(data.get("turns", []))
    directory = Path(directory)
    directory.mkdir(parents=True, exist_ok=True)
    stem = "AI性能评测_" + datetime.now().strftime("%Y%m%d-%H%M%S") + "-" + uuid4().hex[:6]
    paths = {ext: directory / f"{stem}.{ext}" for ext in ("docx", "json")}
    doc = create_document("AI 助手性能评测报告")
    doc.add_paragraph("生成时间：" + data["generated_at"])
    doc.add_paragraph("本报告区分本地知识检索与助手任务表现，在线、离线、降级和数据来源分别统计。")

    def table(rows, widths):
        tab = doc.add_table(rows=0, cols=len(widths))
        tab.style = "Table Grid"
        tab.autofit = False
        for column, width in zip(tab.columns, widths):
            column.width = Cm(width)
        for index, values in enumerate(rows):
            row = tab.add_row()
            for cell, value, width in zip(row.cells, values, widths):
                cell.width = Cm(width)
                cell.text = str(value)
                for p in cell.paragraphs:
                    for r in p.runs:
                        r.font.size = Pt(9)
                        r.bold = index == 0
            # 保持每条短记录完整，跨页重复表头。
            cant_split = OxmlElement("w:cantSplit")
            row._tr.get_or_add_trPr().append(cant_split)
            if index == 0:
                row._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
        return tab

    doc.add_heading("指标汇总", 1)
    table(summary_rows(data), [5.8, 4.0, 7.2])
    doc.add_heading("计算口径", 1)
    for line in (
        "检索 Hit@3 = 前三条结果命中任一预标注正确知识块的题数 / 固定题集总题数。查询失败仍计入分母；标注缺失或失效时不输出整体百分比。",
        "任务完成率 = 验收完成轮次 / 已验收轮次。普通回答先待验收，由实验者按原始请求判定是否全部达到目标；异常请求默认未完成，可复核更正。取消单列且不计入分母。",
        "助手耗时从接收请求到获得最终回答计时，含检索、模型、工具与输出检查；包含异常请求，排除用户取消。不把首 token 到达时间作为最终可用答案时间。",
        "耗时展示中位数与 P95。P95 使用从小到大排序后的第 ceil(0.95×样本数) 个值，单位秒；样本数同时列出。",
        "固定题集按现有正文人工整理，属于开发题集；标注先于运行，不根据本次结果改答案。Hit@3 不等同于回答正确率或物理诊断准确率。",
    ):
        doc.add_paragraph(line)
    run = data.get("retrieval")
    if run:
        doc.add_heading("检索逐题结果", 1)
        doc.add_paragraph(f"题集：{run['version']}；语料块数：{run['corpus_chunks']}；方法：{run['method']}。")
        doc.add_paragraph(run["timing_note"] + f" 索引加载耗时 {run['index_load_s']:.3f} s。")
        rows = [["题号 / 问题", "正确知识块", "实际前 3 条", "命中 / 秒"]]
        for row in run["rows"]:
            result = "标注失效" if row["hit"] is None else "命中" if row["hit"] else "未命中"
            if row["error"]:
                result += "；查询异常"
            rows.append([row["id"] + " " + row["query"], "\n".join(row["expected_ids"]),
                         "\n".join(item["id"] for item in row["retrieved"]) or "无结果",
                         f"{result}\n{row['elapsed_s']:.4f}"])
        table(rows, [6.5, 3.5, 4.3, 2.7])
        doc.add_paragraph("题集 SHA-256：" + run["dataset_sha256"])
        doc.add_paragraph("语料 SHA-256：" + run["corpus_sha256"])
        doc.add_paragraph("正确块正文标注依据、检索标题/页码、错误信息与原始耗时见同名 JSON。")
    doc.add_heading("助手任务逐条验收", 1)
    if data.get("turns"):
        rows = [["问题 / 模式", "验收结果", "耗时 / 验收说明"]]
        for turn in data["turns"]:
            verdict = "取消（不计分）" if turn["status"] == "cancelled" else (
                "待验收" if turn.get("completed") is None else "完成" if turn["completed"] else "未完成")
            scope = MODES[turn["mode"]] + " / " + turn["data_source"]
            rows.append([turn["question"] + "\n" + scope, verdict,
                         seconds_text(turn["elapsed_s"]) + "\n" + turn.get("review_note", "")])
        table(rows, [9.0, 2.5, 5.5])
    else:
        doc.add_paragraph("本次没有已记录的助手交互；任务完成率与助手响应耗时均未评定。")
    if data.get("persistence_error"):
        doc.add_paragraph("会话记录写盘失败：" + data["persistence_error"])
    footer = doc.sections[0].footer.paragraphs[0]
    footer.alignment = 2
    footer.add_run("AI 助手性能评测 · ")
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    try:
        doc.save(paths["docx"])
        paths["json"].write_text(json.dumps(data, ensure_ascii=False, indent=2, allow_nan=False), encoding="utf-8")
    except Exception:
        for path in paths.values():
            path.unlink(missing_ok=True)
        raise
    return {ext: str(path) for ext, path in paths.items()}
