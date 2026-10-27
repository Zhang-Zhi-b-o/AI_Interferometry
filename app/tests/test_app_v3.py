"""v3 本地诊断、出处、冻结记录与 DOCX 的公开路径；无模型付费/真机操作。"""
import json
import threading
import time
from types import SimpleNamespace as NS
from zipfile import ZipFile

from docx import Document
import pytest

from app.agent.evidence import supported_answer
from app.agent.runtime import AgentRuntime
from app.agent.tools import ToolRegistry
from app.core.timeline import Timeline
from app.services.compute import ComputeService
from app.services.diagnostics import diagnose
from app.services.knowledge_index import KnowledgeIndex
from app.services.report import generate_report, prepare_report
from app.engine.agent.provider import ChatResult, ProviderError
from tests.test_app_v2_agent import context


REF = {"id": "S01-p005-01", "source": "S01", "title": "本地测量资料",
       "kind": "paper", "page": 5, "text": "检查干涉条纹的对比度与稳定性。厚度测量应记录原始读数。",
       "path": "资料.pdf"}
INDEX = NS(search=lambda *a, **kw: [REF])


def payload():
    compute = ComputeService()
    rows = [{"d1": 1.0, "d2": d2, "value_mm": compute.thickness_mm(1.0, d2),
             "refractive_index": compute.refractive_index, "event_id": f"E{i:04d}"}
            for i, d2 in enumerate((2.0, 2.001, 1.999), 1)]
    return {"data_source": "demo", "state": {}, "roi": None, "timeline": [],
            "session_id": "test-session", "measurements": rows,
            "uncertainty_parameters": {"micrometer_accuracy_mm": 0.0001,
                "refractive_index_tolerance": 0.001, "confidence": 0.95}}


def test_timeline_bound_copy_persistence_and_threads(tmp_path):
    timeline = Timeline(tmp_path, capacity=5)
    value = {"value": [1]}
    timeline.add("initial", result=value)
    value["value"].append(2)
    assert timeline.snapshot()[0]["result"] == {"value": [1]}
    threads = [threading.Thread(target=lambda: [timeline.add("tick") for _ in range(5)])
               for _ in range(3)]
    for t in threads:
        t.start()
    for t in threads:
        t.join()
    assert len(timeline.snapshot()) == 5
    rows = [json.loads(line) for line in timeline.path.read_text(encoding="utf-8").splitlines()]
    assert len(rows) == 16
    assert len({r["id"] for r in rows}) == 16
    timeline.snapshot()[0]["kind"] = "changed"
    assert timeline.snapshot()[0]["kind"] == "tick"


def test_timeline_disk_error_does_not_block_stop(tmp_path):
    blocked = tmp_path / "file"
    blocked.write_text("x")
    timeline = Timeline(blocked)
    timeline.add("emergency_stop")
    assert timeline.snapshot()[0]["kind"] == "emergency_stop"
    assert timeline.persistence_error


@pytest.mark.parametrize(("state", "extra", "codes"), [
    ({}, {}, {"no_frame"}),
    ({"frame_size": [640, 480], "detecting": True}, {}, {"no_fringe"}),
    ({"frame_size": [640, 480], "center_x": 100, "fringe_width_px": 10,
      "fringe_spacing_px": 5, "tilt_deg": 1.2}, {}, {"off_center", "dense", "tilted"}),
    ({"reading_mm": 1}, {"reading_age_seconds": 2, "reading_stable": False},
     {"no_frame", "reading_stale", "reading_unstable"}),
    ({"frame_size": [640, 480], "center_x": 320, "fringe_width_px": 10},
     {"map_mode": "relative"}, {"uncalibrated_map"}),
])
def test_five_diagnostic_snapshots(state, extra, codes):
    snap = {"state": state, **extra}
    first = diagnose(snap)
    assert first == diagnose(snap)
    assert {f["code"] for f in first["findings"]} == codes


def test_contrast_and_failure_trends_are_relative():
    snap = {"state": {"frame_size": [640, 480]}, "roi": None,
            "timeline": [{"kind": "fringe", "result": {"contrast": c, "roi": None}}
                         for c in (0.8, 0.6, 0.4)] + [
                {"kind": "search_finished", "result": {"failed": True}} for _ in range(2)]}
    assert {f["code"] for f in diagnose(snap)["findings"]} == {"contrast_falling", "search_failures"}
    snap["roi"] = [0, 0, 20, 20]
    assert "contrast_falling" not in {f["code"] for f in diagnose(snap)["findings"]}


def test_docx_records_native_equation_chart_and_references(tmp_path):
    original = payload()
    result = generate_report(original, tmp_path, index=INDEX)
    paths = result["paths"]
    doc = Document(paths["docx"])
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "演示数据" in text
    assert "[K1]" in text and "S01-p005-01" in text
    assert len(doc.tables) == 3
    assert len(doc.inline_shapes) == 1
    with ZipFile(paths["docx"]) as z:
        xml = z.read("word/document.xml").decode()
    assert "<m:oMath" in xml
    assert 'w:val="superscript"' in xml
    data = json.loads(open(paths["json"], encoding="utf-8").read())
    assert data["measurements"] == original["measurements"]
    assert data["statistics"]["mean"] == pytest.approx(ComputeService().thickness_mm(1, 2))
    assert data["uncertainty"]["coverage_factor"] > 1
    assert set(paths) == {"docx", "md", "json"}
    assert "diagnosis" not in original


def test_single_empty_mixed_and_inconsistent_records(tmp_path):
    data = payload()
    data["measurements"] = data["measurements"][:1]
    report = prepare_report(data, index=INDEX)
    assert report["statistics"]["std"] is None
    assert "仅有一次" in report["warnings"][0]
    data["measurements"] = []
    assert prepare_report(data, index=INDEX)["statistics"] == {}
    data = payload()
    row = data["measurements"][1]
    row["refractive_index"] = 1.6
    row["value_mm"] = ComputeService(refractive_index=1.6).thickness_mm(row["d1"], row["d2"])
    assert "折射率不同" in prepare_report(data, index=INDEX)["warnings"][0]
    row["value_mm"] = 999
    with pytest.raises(ValueError, match="不一致"):
        generate_report(data, tmp_path, index=INDEX)
    assert not list(tmp_path.iterdir())


def test_search_custom_chunks_and_source_path(tmp_path):
    index = KnowledgeIndex(chunks=[REF, {**REF, "id": "S02-p001-01", "text": "检查条纹的对比度与稳定性。"},
                                  {**REF, "id": "S03-p001-01", "text": "目录 实验 . . . . . . . . . . . 1"}])
    assert len(index.chunks) == 2
    assert index.search("条纹对比度")[0]["path"] == "资料.pdf"
    assert index.search("完全无关单词abcxyz") == []


def test_output_guard_and_offline_fallback(monkeypatch):
    assert supported_answer("厚度为 0.1230 mm [S01-p005-01]", [{"thickness_mm": 0.123}], [REF])
    assert not supported_answer("厚度为 999 mm", [{"thickness_mm": 0.123}], [REF])
    assert not supported_answer("请调节 [S99-p001-01]", [], [REF])
    monkeypatch.setattr("app.agent.runtime.default_index", lambda: INDEX)
    events = []

    class Bad:
        available = True

        def chat_with_tools(self, *a, on_delta, **kw):
            on_delta("厚度为 999 mm")
            return ChatResult(content="厚度为 999 mm [S99-p001-01]")

    runtime = AgentRuntime(context(), provider=Bad(), on_event=lambda *e: events.append(e))
    assert "厚度为 999 mm" in runtime.ask("下一步")
    assert [e for e in events if e[0] == "warning"]
    assert runtime.last_turn["mode"] == "online"

    class Disconnected:
        available = True

        def chat_with_tools(self, *a, **kw):
            raise ProviderError("offline")

    runtime.provider = Disconnected()
    with pytest.raises(ProviderError, match="没有切换为本地回答"):
        runtime.ask("下一步")
    assert runtime.ctx.calls == []


def test_offline_generate_report_tool(monkeypatch, tmp_path):
    ctx = context()
    ctx.demo_mode = True
    ctx.timeline = Timeline()
    ctx.report_snapshot = payload
    ctx.record = NS(out_dir=tmp_path)
    monkeypatch.setattr("app.services.report.default_index", lambda: INDEX)
    runtime = AgentRuntime(ctx, provider=NS(available=False))
    assert "已生成 Word 报告" in runtime.ask("生成报告")
    assert len(list(tmp_path.glob("*.docx"))) == 1
    assert ctx.timeline.snapshot()[-1]["kind"] == "report_generated"
    assert ctx.calls == []
    registry = ToolRegistry(ctx)
    assert registry.execute("read_timeline", {})
    for query in ("", 4, "a" * 121):
        with pytest.raises(ValueError):
            registry.execute("search_knowledge", {"query": query})


@pytest.mark.parametrize("mode", ["valid", "invented", "disconnected", "cancelled"])
def test_model_report_narrative_and_fallback(tmp_path, mode):
    cancel = threading.Event()

    class Narrator:
        available = True

        def chat_with_tools(self, *args, **kwargs):
            if mode == "disconnected":
                raise ProviderError("offline")
            if mode == "cancelled":
                cancel.set()
            text = "记录两次读数并计算厚度。" if mode != "invented" else "厚度为 999 mm。"
            return ChatResult(content=json.dumps({"purpose": text, "steps": text,
                              "analysis": "核对条纹稳定性。[S01-p005-01]", "conclusion": text}))

    from app.engine.agent.provider import ProviderCancelled
    if mode == "cancelled":
        with pytest.raises(ProviderCancelled):
            generate_report(payload(), tmp_path, index=INDEX, provider=Narrator(), cancel_event=cancel)
        assert not list(tmp_path.iterdir())
        return
    result = generate_report(payload(), tmp_path, index=INDEX, provider=Narrator(), cancel_event=cancel)
    data = json.loads(open(result["paths"]["json"], encoding="utf-8").read())
    assert bool(data["narrative"]) == (mode == "valid")
    assert "999" not in "\n".join(p.text for p in Document(result["paths"]["docx"]).paragraphs)


def test_cancel_discards_buffered_deltas(monkeypatch):
    monkeypatch.setattr("app.agent.runtime.default_index", lambda: INDEX)
    entered, release = threading.Event(), threading.Event()
    events = []

    class Delayed:
        available = True

        def chat_with_tools(self, *a, on_delta, **kw):
            on_delta("迟到草稿")
            entered.set()
            release.wait(3)
            return ChatResult(content="迟到草稿")

    runtime = AgentRuntime(context(), provider=Delayed(), on_event=lambda *e: events.append(e))
    thread = threading.Thread(target=lambda: runtime.ask("下一步"))
    thread.start()
    assert entered.wait(2)
    runtime.cancel()
    thread.join(1)
    delivered_before_release = [e for e in events if e[0] == "delta"]
    release.set()
    assert not thread.is_alive()
    time.sleep(0.02)
    assert [e for e in events if e[0] == "delta"] == delivered_before_release


def test_missing_index_still_gives_offline_diagnosis(monkeypatch):
    def unavailable():
        raise OSError("index unavailable")

    monkeypatch.setattr("app.agent.runtime.default_index", unavailable)
    ctx = context()
    ctx.demo_mode = True
    runtime = AgentRuntime(ctx, provider=NS(available=False))
    answer = runtime.ask("下一步")
    assert "离线模式" in answer and "观察" in answer
    assert runtime.ctx.calls == []
